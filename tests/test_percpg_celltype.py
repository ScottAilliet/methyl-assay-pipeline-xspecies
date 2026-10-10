"""Test per-CpG per-cell-type beta value enrichment.

Verifies that:
1. _build_all_beta_group_map correctly maps beta files to atlas groups
2. enrich_percpg_per_celltype correctly computes per-CpG per-cell-type means
3. enrich_blocks_with_celltype_beta (JSON dict version) produces identical results
4. Low-coverage CpGs are excluded (cov < min_cov -> NaN)
5. Sheet 4 "CpG CellType Beta" is generated with correct structure
"""
import json
import os
import tempfile

import numpy as np
import pytest

from methyl_panel.phase1_dmr_loader import CpGSite, DMRBlock
from methyl_panel.phase0_dmr_discovery import (
    enrich_percpg_per_celltype,
    enrich_blocks_with_celltype_beta,
    _build_all_beta_group_map,
)
from methyl_panel.phase3_primer3_design import PrimerPair


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _create_synthetic_beta_file(path, cpg_data):
    """Create a synthetic .beta file.

    Args:
        path: Output file path
        cpg_data: dict mapping 1-based CpG index -> (meth, cov) tuple.
                  The file spans CpG indices 1..max(cpg_data.keys()).
    """
    max_idx = max(cpg_data.keys())
    arr = np.zeros((max_idx, 2), dtype=np.uint8)
    for idx, (meth, cov) in cpg_data.items():
        arr[idx - 1] = [meth, cov]
    arr.tofile(path)


def _make_block(cpg_indices, cell_type_id="TypeA"):
    """Create a DMRBlock with CpGSites at the given global indices."""
    cpg_sites = []
    for j, gidx in enumerate(cpg_indices):
        cpg_sites.append(CpGSite(
            label=f"CpG_{j+1}",
            position=1000 + j * 50,
            global_idx=gidx,
            target_mean_beta=0.0,
            background_mean_beta=0.0,
            delta_beta=0.0,
        ))
    return DMRBlock(
        cell_type_id=cell_type_id, seq_id=f"{cell_type_id}_0001", rank=1,
        chrom="chr1", start=1000, end=1100, num_cpgs=len(cpg_indices),
        block_len=100, gene="TEST", annotation="test", direction="hypo",
        target_mean_meth=0.25, background_mean_meth=0.5,
        delta_means=0.25, delta_quants=0.3, delta_maxmin=0.7,
        ttest_pval="0.01", mwtest_pval="0.02", mvalue_ttest="0.03",
        dms_score=0.8, target_celltype=cell_type_id,
        cpg_sites=cpg_sites,
    )


# Beta values for the synthetic data:
#   CpG 100: TypeA=[0.8, 0.7], TypeB=[0.1, 0.2], TypeC=[0.5, 0.5]
#   CpG 101: TypeA=[0.2, 0.3], TypeB=[0.9, 0.8], TypeC=[0.5, 0.5]
#   CpG 102: TypeA=[0.9, 0.8], TypeB=[0.2, 0.1], TypeC=[0.5, 0.5]
# Expected means:
#   CpG 100: TypeA=0.75, TypeB=0.15, TypeC=0.50
#   CpG 101: TypeA=0.25, TypeB=0.85, TypeC=0.50
#   CpG 102: TypeA=0.85, TypeB=0.15, TypeC=0.50
_BETA_DATA = {
    "TypeA_s1": {100: (8, 10), 101: (2, 10), 102: (9, 10)},
    "TypeA_s2": {100: (7, 10), 101: (3, 10), 102: (8, 10)},
    "TypeB_s1": {100: (1, 10), 101: (9, 10), 102: (2, 10)},
    "TypeB_s2": {100: (2, 10), 101: (8, 10), 102: (1, 10)},
    "TypeC_s1": {100: (5, 10), 101: (5, 10), 102: (5, 10)},
    "TypeC_s2": {100: (5, 10), 101: (5, 10), 102: (5, 10)},
}

_EXPECTED = {
    100: {"TypeA": 0.75, "TypeB": 0.15, "TypeC": 0.50},
    101: {"TypeA": 0.25, "TypeB": 0.85, "TypeC": 0.50},
    102: {"TypeA": 0.85, "TypeB": 0.15, "TypeC": 0.50},
}


@pytest.fixture
def beta_setup():
    """Create synthetic beta files, atlas groups CSV, and beta_group_map."""
    with tempfile.TemporaryDirectory() as tmpdir:
        groups_csv = os.path.join(tmpdir, "atlas_groups.csv")
        with open(groups_csv, "w") as f:
            f.write("name,group\n")
            for sample in _BETA_DATA:
                group = sample.rsplit("_", 1)[0]
                f.write(f"{sample},{group}\n")

        beta_dir = os.path.join(tmpdir, "beta_files")
        os.makedirs(beta_dir)
        for sample, cpg_data in _BETA_DATA.items():
            _create_synthetic_beta_file(
                os.path.join(beta_dir, f"{sample}.beta"), cpg_data
            )

        beta_group_map = _build_all_beta_group_map(groups_csv, beta_dir)
        yield tmpdir, beta_dir, groups_csv, beta_group_map


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------
class TestBuildAllBetaGroupMap:
    def test_correct_count(self, beta_setup):
        _, _, _, beta_group_map = beta_setup
        assert len(beta_group_map) == 6

    def test_correct_groups(self, beta_setup):
        _, _, _, beta_group_map = beta_setup
        groups = set(beta_group_map.values())
        assert groups == {"TypeA", "TypeB", "TypeC"}

    def test_two_samples_per_group(self, beta_setup):
        from collections import Counter
        _, _, _, beta_group_map = beta_setup
        counts = Counter(beta_group_map.values())
        assert counts == {"TypeA": 2, "TypeB": 2, "TypeC": 2}


class TestEnrichPercpgPerCelltype:
    def test_all_cpgs_all_groups(self, beta_setup):
        """Per-CpG per-cell-type means match expected values."""
        _, _, _, beta_group_map = beta_setup
        block = _make_block([100, 101, 102])
        enrich_percpg_per_celltype([block], beta_group_map, min_cov=5)

        for j, gidx in enumerate([100, 101, 102]):
            cpg = block.cpg_sites[j]
            for group, expected in _EXPECTED[gidx].items():
                assert group in cpg.per_celltype_beta, (
                    f"CpG {gidx} missing group {group}"
                )
                assert abs(cpg.per_celltype_beta[group] - expected) < 0.01, (
                    f"CpG {gidx} {group}: got {cpg.per_celltype_beta[group]}, "
                    f"expected {expected}"
                )

    def test_low_coverage_excluded(self, beta_setup):
        """CpGs with coverage < min_cov are excluded from the mean."""
        tmpdir, beta_dir, groups_csv, beta_group_map = beta_setup

        # Overwrite TypeA_s1 with low coverage at CpG 100
        low_cov_data = {100: (8, 3), 101: (2, 10), 102: (9, 10)}  # cov=3 < 5
        _create_synthetic_beta_file(
            os.path.join(beta_dir, "TypeA_s1.beta"), low_cov_data
        )
        beta_group_map = _build_all_beta_group_map(groups_csv, beta_dir)

        block = _make_block([100, 101, 102])
        enrich_percpg_per_celltype([block], beta_group_map, min_cov=5)

        # TypeA CpG100: only sample2 has valid coverage (beta=0.7)
        cpg1 = block.cpg_sites[0]
        assert abs(cpg1.per_celltype_beta["TypeA"] - 0.7) < 0.01

        # TypeA CpG101: both samples valid (0.2 + 0.3) / 2 = 0.25
        cpg2 = block.cpg_sites[1]
        assert abs(cpg2.per_celltype_beta["TypeA"] - 0.25) < 0.01

    def test_empty_blocks(self, beta_setup):
        """Empty blocks list is handled gracefully."""
        _, _, _, beta_group_map = beta_setup
        enrich_percpg_per_celltype([], beta_group_map)
        # No crash, no error

    def test_empty_group_map(self, beta_setup):
        """Empty beta_group_map is handled gracefully."""
        block = _make_block([100, 101, 102])
        enrich_percpg_per_celltype([block], {})
        assert block.cpg_sites[0].per_celltype_beta == {}


class TestEnrichBlocksJsonDict:
    def test_json_dict_matches_dataclass(self, beta_setup):
        """JSON dict enrichment produces same values as DMRBlock enrichment."""
        _, beta_dir, groups_csv, beta_group_map = beta_setup

        # DMRBlock version
        block_obj = _make_block([100, 101, 102])
        enrich_percpg_per_celltype([block_obj], beta_group_map, min_cov=5)

        # JSON dict version
        dmr_blocks = [{
            "seq_id": "TypeA_0001", "chrom": "chr1",
            "start": 1000, "end": 1100, "num_cpgs": 3,
            "cpg_sites": [
                {"label": "CpG_1", "position": 1000, "global_idx": 100,
                 "target_mean_beta": 0.0, "background_mean_beta": 0.0, "delta_beta": 0.0},
                {"label": "CpG_2", "position": 1050, "global_idx": 101,
                 "target_mean_beta": 0.0, "background_mean_beta": 0.0, "delta_beta": 0.0},
                {"label": "CpG_3", "position": 1100, "global_idx": 102,
                 "target_mean_beta": 0.0, "background_mean_beta": 0.0, "delta_beta": 0.0},
            ],
        }]
        enriched = enrich_blocks_with_celltype_beta(dmr_blocks, beta_dir, groups_csv)
        assert enriched == 1

        # Compare values
        for j in range(3):
            obj_pct = block_obj.cpg_sites[j].per_celltype_beta
            dict_pct = dmr_blocks[0]["cpg_sites"][j].get("per_celltype_beta", {})
            assert set(obj_pct.keys()) == set(dict_pct.keys())
            for k in obj_pct:
                assert abs(obj_pct[k] - dict_pct[k]) < 0.001


class TestSheet4Generation:
    def test_sheet4_created_with_celltype_data(self, beta_setup):
        """Sheet 4 is created when per_celltype_beta data is present."""
        
        from methyl_panel.output_xlsx import write_xlsx
        import openpyxl

        tmpdir = beta_setup[0]
        pairs = PrimerPair.load_primers("tests/fixtures/MONO/primers.json")
        with open("tests/fixtures/MONO/dmr_blocks.json") as f:
            blocks = json.load(f)

        # Inject per_celltype_beta into the block matching the first primer
        target_seq_id = pairs[0].seq_id
        target_block = next(b for b in blocks if b["seq_id"] == target_seq_id)
        for c in target_block.get("cpg_sites", []):
            c["per_celltype_beta"] = {"TypeA": 0.75, "TypeB": 0.15, "TypeC": 0.50}

        out_path = os.path.join(tmpdir, "test_sheet4.xlsx")
        write_xlsx(pairs[:5], out_path, dmr_blocks=blocks)

        wb = openpyxl.load_workbook(out_path, read_only=True)
        assert "CpG CellType Beta" in wb.sheetnames
        ws4 = wb["CpG CellType Beta"]
        headers = [c.value for c in next(ws4.iter_rows(min_row=1, max_row=1))]

        # Meta columns present
        for meta in ["assay_id", "seq_id", "cpg_label", "cpg_position"]:
            assert meta in headers

        # Cell type columns present and sorted
        ct_cols = [h for h in headers if h not in
                   {"assay_id", "seq_id", "cpg_label", "cpg_position"}]
        assert "TypeA" in ct_cols
        assert "TypeB" in ct_cols
        assert "TypeC" in ct_cols
        assert ct_cols == sorted(ct_cols)

        # Row count: each primer matching the target block contributes
        # its CpGs, so total rows = n_primers_with_seq_id * n_cpgs
        n_cpgs = len(target_block.get("cpg_sites", []))
        n_matching_primers = sum(1 for p in pairs[:5] if p.seq_id == target_seq_id)
        assert ws4.max_row - 1 == n_cpgs * n_matching_primers
        wb.close()

    def test_sheet4_absent_without_celltype_data(self):
        """Sheet 4 is absent when no per_celltype_beta data exists."""
        
        from methyl_panel.output_xlsx import write_xlsx
        import openpyxl

        with tempfile.TemporaryDirectory() as tmpdir:
            pairs = PrimerPair.load_primers("tests/fixtures/MONO/primers.json")
            with open("tests/fixtures/MONO/dmr_blocks.json") as f:
                blocks = json.load(f)

            out_path = os.path.join(tmpdir, "test_no_sheet4.xlsx")
            write_xlsx(pairs[:5], out_path, dmr_blocks=blocks)

            wb = openpyxl.load_workbook(out_path, read_only=True)
            assert "CpG CellType Beta" not in wb.sheetnames
            assert "DMR CpGs" in wb.sheetnames
            wb.close()


class TestSpecificityMetric:
    """Test the min_non_target_delta / worst_celltype columns in Sheet 1."""

    def _load_monO_data(self):
        """Load MONO primers and blocks, inject per_celltype_beta."""
        

        pairs = PrimerPair.load_primers("tests/fixtures/MONO/primers.json")
        with open("tests/fixtures/MONO/dmr_blocks.json") as f:
            blocks = json.load(f)

        target_seq_id = pairs[0].seq_id
        target_block = next(b for b in blocks if b["seq_id"] == target_seq_id)

        # Inject per_celltype_beta into all CpGs of the target block.
        # Blood-Monocytes is the target (MONO -> Blood-Monocytes).
        # Set target high, others low with varying deltas.
        for c in target_block.get("cpg_sites", []):
            c["per_celltype_beta"] = {
                "Blood-Monocytes": 0.90,  # target (excluded from non_target)
                "Blood-T-CD4": 0.05,      # delta = 0.85
                "Blood-B": 0.0,           # delta = 0.90
                "Adipocytes": 0.10,       # delta = 0.80 (worst)
            }
            c["target_mean_beta"] = 0.90
            c["background_mean_beta"] = 0.05

        return pairs, blocks, target_seq_id

    def test_blood_type_excludes_target_groups(self):
        """MONO specificity metric excludes Blood-Monocytes (target group)."""
        from methyl_panel.output_xlsx import write_xlsx
        import openpyxl

        pairs, blocks, target_seq_id = self._load_monO_data()

        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = os.path.join(tmpdir, "test_specificity_blood.xlsx")
            write_xlsx(pairs[:5], out_path, dmr_blocks=blocks)

            wb = openpyxl.load_workbook(out_path, read_only=True)
            ws1 = wb["Primer pairs"]
            headers = [c.value for c in next(ws1.iter_rows(min_row=1, max_row=1))]

            assert "min_non_target_delta" in headers
            assert "worst_celltype" in headers

            delta_idx = headers.index("min_non_target_delta")
            worst_idx = headers.index("worst_celltype")
            seq_id_idx = headers.index("seq_id")

            found = False
            for row in ws1.iter_rows(min_row=2, values_only=True):
                if row[seq_id_idx] == target_seq_id:
                    # min_non_target_delta = 0.80 (Adipocytes is closest non-target)
                    assert row[delta_idx] is not None, \
                        "min_non_target_delta should be computed for MONO"
                    assert abs(row[delta_idx] - 0.80) < 0.01, \
                        f"Expected 0.80, got {row[delta_idx]}"
                    # worst_celltype = Adipocytes (smallest delta)
                    assert row[worst_idx] == "Adipocytes", \
                        f"Expected Adipocytes, got {row[worst_idx]}"
                    # Target group must NOT appear as worst_celltype
                    assert row[worst_idx] != "Blood-Monocytes"
                    found = True
                    break

            assert found, "No row matching target seq_id found in Sheet 1"
            wb.close()

    def test_atlas_type_direct_match(self):
        """Atlas type (Adipocytes) uses direct name match, no mapping needed."""
        from methyl_panel.output_xlsx import write_xlsx
        import openpyxl

        pairs, blocks, target_seq_id = self._load_monO_data()

        # Override cell_type_id to an atlas type name
        for p in pairs[:5]:
            p.cell_type_id = "Adipocytes"

        # Re-inject per_celltype_beta with Adipocytes as target
        target_block = next(b for b in blocks if b["seq_id"] == target_seq_id)
        for c in target_block.get("cpg_sites", []):
            c["per_celltype_beta"] = {
                "Adipocytes": 0.85,        # target (direct match)
                "Blood-Monocytes": 0.20,   # delta = 0.65
                "Blood-T-CD4": 0.05,       # delta = 0.80
                "Hepatocytes": 0.15,       # delta = 0.70
            }
            c["target_mean_beta"] = 0.85
            c["background_mean_beta"] = 0.20

        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = os.path.join(tmpdir, "test_specificity_atlas.xlsx")
            write_xlsx(pairs[:5], out_path, dmr_blocks=blocks)

            wb = openpyxl.load_workbook(out_path, read_only=True)
            ws1 = wb["Primer pairs"]
            headers = [c.value for c in next(ws1.iter_rows(min_row=1, max_row=1))]

            delta_idx = headers.index("min_non_target_delta")
            worst_idx = headers.index("worst_celltype")
            seq_id_idx = headers.index("seq_id")

            found = False
            for row in ws1.iter_rows(min_row=2, values_only=True):
                if row[seq_id_idx] == target_seq_id:
                    # Deltas: Blood-Monocytes=0.65, Hepatocytes=0.70, Blood-T-CD4=0.80
                    # Worst = Blood-Monocytes (0.65)
                    assert row[delta_idx] is not None
                    assert abs(row[delta_idx] - 0.65) < 0.01, \
                        f"Expected 0.65, got {row[delta_idx]}"
                    assert row[worst_idx] == "Blood-Monocytes", \
                        f"Expected Blood-Monocytes, got {row[worst_idx]}"
                    # Adipocytes (target) must NOT be worst
                    assert row[worst_idx] != "Adipocytes"
                    found = True
                    break

            assert found, "No row matching target seq_id found in Sheet 1"
            wb.close()

    def test_no_celltype_data_returns_none(self):
        """Specificity metric is None when no per_celltype_beta data exists."""
        
        from methyl_panel.output_xlsx import write_xlsx
        import openpyxl

        with tempfile.TemporaryDirectory() as tmpdir:
            pairs = PrimerPair.load_primers("tests/fixtures/MONO/primers.json")
            with open("tests/fixtures/MONO/dmr_blocks.json") as f:
                blocks = json.load(f)

            out_path = os.path.join(tmpdir, "test_no_specificity.xlsx")
            write_xlsx(pairs[:5], out_path, dmr_blocks=blocks)

            wb = openpyxl.load_workbook(out_path, read_only=True)
            ws1 = wb["Primer pairs"]
            headers = [c.value for c in next(ws1.iter_rows(min_row=1, max_row=1))]

            delta_idx = headers.index("min_non_target_delta")
            worst_idx = headers.index("worst_celltype")

            # All rows should have None for specificity (no per_celltype_beta)
            for row in ws1.iter_rows(min_row=2, values_only=True):
                assert row[delta_idx] is None
                assert row[worst_idx] is None
                break  # Just check first data row
            wb.close()


class TestIsBloodCellType:
    """Test that is_blood_cell_type only returns True for merged blood types."""

    def test_merged_blood_types_return_true(self):
        """The 7 merged blood types (MONO, BCELL, NK, GRAN, CD3T, CD4T, CD8T)
        must return True — they need blood-only background + per-subgroup filter."""
        from methyl_panel.phase0_dmr_discovery import is_blood_cell_type

        for ct in ["MONO", "BCELL", "NK", "GRAN", "CD3T", "CD4T", "CD8T"]:
            assert is_blood_cell_type(ct) is True, \
                f"{ct} should be a blood cell type"

    def test_atlas_blood_types_return_false(self):
        """Individual atlas blood groups (Blood-T-CD3, Blood-Monocytes, etc.)
        must return False — they should be treated like any other atlas type
        (full atlas background, per-subgroup filter disabled)."""
        from methyl_panel.phase0_dmr_discovery import is_blood_cell_type

        atlas_blood = [
            "Blood-T-CD3", "Blood-T-CD4", "Blood-T-CD8",
            "Blood-Monocytes", "Blood-B", "Blood-B-Mem",
            "Blood-NK", "Blood-Granulocytes",
            "Blood-T-CenMem-CD4", "Blood-T-Eff-CD8",
            "Blood-T-EffMem-CD4", "Blood-T-EffMem-CD8",
            "Blood-T-Naive-CD4", "Blood-T-Naive-CD8",
        ]
        for ct in atlas_blood:
            assert is_blood_cell_type(ct) is False, \
                f"{ct} should NOT be treated as a merged blood cell type"

    def test_non_blood_types_return_false(self):
        """Non-blood atlas types must return False."""
        from methyl_panel.phase0_dmr_discovery import is_blood_cell_type

        for ct in ["Adipocytes", "Liver-Hepatocytes", "Cortex-Neuron"]:
            assert is_blood_cell_type(ct) is False
